import os
import sys
import tempfile
import threading
import csv
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import customtkinter as ctk

from classroom_manager import ClassroomManager
from sw_analyzer import analyze_file_submission
from duplicate_detector import analyze_duplicates

# Opciones globales de apariencia
ctk.set_appearance_mode("System")  # "System", "Dark", "Light"
ctk.set_default_color_theme("blue")

class SolidWorksClassroomApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Inspector de Autores SolidWorks - Google Classroom")
        self.geometry("1100 HIGHx720".replace("HIGHx", "x"))
        self.minsize(950, 650)

        # Gestores de backend
        self.classroom_mgr = ClassroomManager()
        self.courses_data = []
        self.assignments_data = []
        self.analysis_rows = []
        self.stats_data = {}

        # Construir Interfaz de Usuario
        self._create_widgets()
        
        # Verificar estado de inicio de sesión inicial
        self.after(500, self._check_initial_auth)

    def _create_widgets(self):
        # 1. Header Frame (Autenticación e Identidad)
        self.header_frame = ctk.CTkFrame(self, corner_radius=10)
        self.header_frame.pack(fill="x", padx=15, pady=(15, 10))

        self.title_label = ctk.CTkLabel(
            self.header_frame,
            text="🛠️ Inspector de Autores de SolidWorks en Classroom",
            font=ctk.CTkFont(size=18, weight="bold")
        )
        self.title_label.pack(side="left", padx=15, pady=12)

        self.auth_status_label = ctk.CTkLabel(
            self.header_frame,
            text="Estado: Verificando...",
            font=ctk.CTkFont(size=13)
        )
        self.auth_status_label.pack(side="right", padx=15, pady=12)

        self.login_btn = ctk.CTkButton(
            self.header_frame,
            text="Iniciar Sesión",
            width=120,
            command=self._handle_login
        )
        self.login_btn.pack(side="right", padx=5, pady=12)

        # 2. Control Panel Frame (Selección de Clase y Tarea)
        self.control_frame = ctk.CTkFrame(self, corner_radius=10)
        self.control_frame.pack(fill="x", padx=15, pady=5)

        # Fila 1: Selección de Clase
        self.lbl_class = ctk.CTkLabel(self.control_frame, text="1. Seleccionar Clase:", font=ctk.CTkFont(size=13, weight="bold"))
        self.lbl_class.grid(row=0, column=0, padx=15, pady=10, sticky="w")

        self.combo_courses = ctk.CTkOptionMenu(self.control_frame, width=380, values=["Cargando clases..."])
        self.combo_courses.grid(row=0, column=1, padx=10, pady=10, sticky="w")

        self.btn_load_courses = ctk.CTkButton(self.control_frame, text="🔄 Recargar Clases", width=120, command=self._load_courses_async)
        self.btn_load_courses.grid(row=0, column=2, padx=10, pady=10)

        # Fila 2: Selección de Tarea
        self.lbl_task = ctk.CTkLabel(self.control_frame, text="2. Seleccionar Tarea:", font=ctk.CTkFont(size=13, weight="bold"))
        self.lbl_task.grid(row=1, column=0, padx=15, pady=10, sticky="w")

        self.combo_assignments = ctk.CTkOptionMenu(self.control_frame, width=380, values=["Selecciona una clase primero"])
        self.combo_assignments.grid(row=1, column=1, padx=10, pady=10, sticky="w")

        self.btn_load_tasks = ctk.CTkButton(self.control_frame, text="📋 Cargar Tareas", width=120, command=self._load_assignments_async)
        self.btn_load_tasks.grid(row=1, column=2, padx=10, pady=10)

        # Fila 3: Botón Principal de Acción
        self.btn_analyze = ctk.CTkButton(
            self.control_frame,
            text="🚀 Descargar y Analizar Piezas de SolidWorks",
            font=ctk.CTkFont(size=14, weight="bold"),
            fg_color="#2b8a3e",
            hover_color="#216a30",
            height=40,
            command=self._start_analysis_async
        )
        self.btn_analyze.grid(row=2, column=0, columnspan=3, padx=15, pady=(5, 15), sticky="ew")

        # Configurar expansión de columnas
        self.control_frame.grid_columnconfigure(1, weight=1)

        # 3. Progress & Status Frame
        self.progress_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.progress_frame.pack(fill="x", padx=15, pady=5)

        self.progress_bar = ctk.CTkProgressBar(self.progress_frame)
        self.progress_bar.pack(fill="x", pady=2)
        self.progress_bar.set(0)

        self.status_label = ctk.CTkLabel(self.progress_frame, text="Listo para analizar.", font=ctk.CTkFont(size=12))
        self.status_label.pack(anchor="w", pady=2)

        # 4. Stat Summary Cards Frame
        self.stats_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.stats_frame.pack(fill="x", padx=15, pady=5)

        self.card_total_students = self._create_stat_card(self.stats_frame, "Total Alumnos", "0", 0)
        self.card_total_files = self._create_stat_card(self.stats_frame, "Archivos Analyzados", "0", 1)
        self.card_suspicious = self._create_stat_card(self.stats_frame, "⚠️ Creadores Repetidos", "0", 2, color="#c92a2a")

        # 5. Filter & Export Bar
        self.filter_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.filter_frame.pack(fill="x", padx=15, pady=(5, 0))

        self.chk_only_duplicates = ctk.CTkCheckBox(
            self.filter_frame,
            text="Mostrar únicamente creadores repetidos / sospechosos",
            command=self._apply_table_filter
        )
        self.chk_only_duplicates.pack(side="left", padx=5)

        self.btn_export_excel = ctk.CTkButton(
            self.filter_frame,
            text="📥 Exportar Reporte a Excel",
            width=180,
            fg_color="#1864ab",
            hover_color="#10497e",
            command=self._export_to_excel
        )
        self.btn_export_excel.pack(side="right", padx=5)

        # 6. Result Table (Treeview)
        self.table_frame = ctk.CTkFrame(self, corner_radius=10)
        self.table_frame.pack(fill="both", expand=True, padx=15, pady=10)

        columns = ("status", "student", "file_name", "sw_author", "sw_last_author", "details")
        self.tree = ttk.Treeview(self.table_frame, columns=columns, show="headings", selectmode="browse")

        self.tree.heading("status", text="Estado / Alerta")
        self.tree.heading("student", text="Alumno (Classroom)")
        self.tree.heading("file_name", text="Nombre de la Pieza / Archivo")
        self.tree.heading("sw_author", text="Nombre grabado en SW (Creador)")
        self.tree.heading("sw_last_author", text="Último cambio por (SW)")
        self.tree.heading("details", text="Detalles de Coincidencia")

        self.tree.column("status", width=140, anchor="center")
        self.tree.column("student", width=180, anchor="w")
        self.tree.column("file_name", width=220, anchor="w")
        self.tree.column("sw_author", width=180, anchor="w")
        self.tree.column("sw_last_author", width=160, anchor="w")
        self.tree.column("details", width=250, anchor="w")

        # Scrollbars para el Treeview
        vsb = ttk.Scrollbar(self.table_frame, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(self.table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

        self.table_frame.grid_rowconfigure(0, weight=1)
        self.table_frame.grid_columnconfigure(0, weight=1)

        # Estilo visual de la tabla
        self._style_treeview()

    def _create_stat_card(self, parent, title, value, col, color=None):
        card = ctk.CTkFrame(parent, corner_radius=8)
        card.grid(row=0, column=col, padx=5, pady=5, sticky="ew")
        parent.grid_columnconfigure(col, weight=1)

        lbl_t = ctk.CTkLabel(card, text=title, font=ctk.CTkFont(size=12))
        lbl_t.pack(anchor="w", padx=10, pady=(8, 2))

        lbl_v = ctk.CTkLabel(
            card,
            text=value,
            font=ctk.CTkFont(size=20, weight="bold"),
            text_color=color if color else ("#1a1a1a", "#ffffff")
        )
        lbl_v.pack(anchor="w", padx=10, pady=(0, 8))

        return lbl_v

    def _style_treeview(self):
        style = ttk.Style()
        style.theme_use("clam")
        
        # Configurar colores de filas y encabezados
        style.configure(
            "Treeview",
            background="#2b2b2b" if ctk.get_appearance_mode() == "Dark" else "#ffffff",
            foreground="#ffffff" if ctk.get_appearance_mode() == "Dark" else "#000000",
            fieldbackground="#2b2b2b" if ctk.get_appearance_mode() == "Dark" else "#ffffff",
            rowheight=26,
            font=('Calibri', 10)
        )
        style.configure(
            "Treeview.Heading",
            background="#1f1f1f" if ctk.get_appearance_mode() == "Dark" else "#e0e0e0",
            foreground="#ffffff" if ctk.get_appearance_mode() == "Dark" else "#000000",
            font=('Calibri', 10, 'bold')
        )
        
        # Tags de filas para destacar duplicados
        self.tree.tag_configure("ALERT", background="#ffc9c9", foreground="#900000")
        self.tree.tag_configure("ALERT_DARK", background="#5c0000", foreground="#ffb3b3")
        self.tree.tag_configure("NORMAL", background="#ffffff" if ctk.get_appearance_mode() == "Light" else "#2b2b2b")

    def _check_initial_auth(self):
        """Verifica la sesión guardada en token.json sin bloquear la UI."""
        if self.classroom_mgr.is_authenticated():
            self._handle_login(silent=True)
        else:
            self.auth_status_label.configure(text="Sesión: No autenticado", text_color="#e03131")
            self.login_btn.configure(text="Iniciar Sesión")
            self.combo_courses.configure(values=["Inicia sesión para cargar clases"])

    def _handle_login(self, silent=False):
        """Realiza el proceso de login en un hilo secundario para no congelar la ventana."""
        def auth_thread():
            try:
                email = self.classroom_mgr.authenticate()
                self.after(0, lambda: self._on_login_success(email))
            except Exception as e:
                if not silent:
                    self.after(0, lambda: messagebox.showerror("Error de Autenticación", f"No se pudo iniciar sesión:\n{e}"))
                self.after(0, lambda: self.auth_status_label.configure(text="Error de Sesión", text_color="#e03131"))

        self.auth_status_label.configure(text="Autenticando en navegador...", text_color="#f59f00")
        threading.Thread(target=auth_thread, daemon=True).start()

    def _on_login_success(self, email):
        self.auth_status_label.configure(text=f"Conectado: {email}", text_color="#2b8a3e")
        self.login_btn.configure(text="Cambiar Cuenta")
        self._load_courses_async()

    def _load_courses_async(self):
        """Carga las clases de Google Classroom asíncronamente."""
        self.status_label.configure(text="Cargando clases desde Google Classroom...")
        self.progress_bar.set(0.2)
        
        def run():
            try:
                courses = self.classroom_mgr.list_courses()
                self.courses_data = courses
                course_names = [f"{c['name']} ({c['section']})" if c['section'] else c['name'] for c in courses]
                if not course_names:
                    course_names = ["No se encontraron clases activas"]
                
                def update_ui():
                    self.combo_courses.configure(values=course_names)
                    if course_names:
                        self.combo_courses.set(course_names[0])
                        self._load_assignments_async()
                    self.status_label.configure(text="Clases cargadas correctamente.")
                    self.progress_bar.set(0)

                self.after(0, update_ui)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Error al cargar clases:\n{e}"))
                self.after(0, lambda: self.status_label.configure(text="Error al cargar clases."))

        threading.Thread(target=run, daemon=True).start()

    def _load_assignments_async(self):
        """Carga las tareas de la clase seleccionada."""
        selected_index = self.combo_courses.cget("values").index(self.combo_courses.get()) if self.combo_courses.get() in self.combo_courses.cget("values") else -1
        if selected_index < 0 or selected_index >= len(self.courses_data):
            return

        course_id = self.courses_data[selected_index]['id']
        self.status_label.configure(text="Cargando tareas de la clase...")
        self.progress_bar.set(0.4)

        def run():
            try:
                assignments = self.classroom_mgr.list_assignments(course_id)
                self.assignments_data = assignments
                task_names = [a['title'] for a in assignments]
                if not task_names:
                    task_names = ["No hay tareas en esta clase"]

                def update_ui():
                    self.combo_assignments.configure(values=task_names)
                    if task_names:
                        self.combo_assignments.set(task_names[0])
                    self.status_label.configure(text="Tareas cargadas.")
                    self.progress_bar.set(0)

                self.after(0, update_ui)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Error al obtener tareas:\n{e}"))
                self.after(0, lambda: self.status_label.configure(text="Error al cargar tareas."))

        threading.Thread(target=run, daemon=True).start()

    def _start_analysis_async(self):
        """Descarga las entregas de la tarea y analiza las piezas de SolidWorks."""
        course_idx = self.combo_courses.cget("values").index(self.combo_courses.get()) if self.combo_courses.get() in self.combo_courses.cget("values") else -1
        task_idx = self.combo_assignments.cget("values").index(self.combo_assignments.get()) if self.combo_assignments.get() in self.combo_assignments.cget("values") else -1

        if course_idx < 0 or course_idx >= len(self.courses_data):
            messagebox.showwarning("Selección Inválida", "Por favor selecciona una clase válida.")
            return

        if task_idx < 0 or task_idx >= len(self.assignments_data):
            messagebox.showwarning("Selección Inválida", "Por favor selecciona una tarea válida.")
            return

        course_id = self.courses_data[course_idx]['id']
        coursework_id = self.assignments_data[task_idx]['id']

        self.btn_analyze.configure(state="disabled")
        self.status_label.configure(text="Iniciando descarga temporal de entregas...")
        self.progress_bar.set(0.1)

        def progress_update(msg, pct):
            self.after(0, lambda: self.status_label.configure(text=msg))
            self.after(0, lambda: self.progress_bar.set(pct))

        def run_pipeline():
            temp_dir = tempfile.mkdtemp(prefix="sw_classroom_")
            try:
                # 1. Descargar entregas de Classroom a carpeta temporal
                downloaded_files = self.classroom_mgr.download_submissions_files(
                    course_id, coursework_id, temp_dir, progress_callback=progress_update
                )

                if not downloaded_files:
                    self.after(0, lambda: messagebox.showinfo("Sin Archivos", "No se encontraron archivos de SolidWorks (.sldprt, .sldasm, .zip) adjuntos en esta tarea."))
                    self.after(0, lambda: self._reset_analysis_ui())
                    return

                # 2. Analizar metadatos de SolidWorks
                progress_update("Analizando metadatos internos de SolidWorks (.sldprt / .sldasm)...", 0.8)
                all_sw_results = []
                for item in downloaded_files:
                    sw_pieces = analyze_file_submission(
                        item['local_path'], item['student_name'], item['student_email'], temp_dir
                    )
                    all_sw_results.extend(sw_pieces)

                # 3. Detección de duplicados / plagio
                progress_update("Evaluando coincidencias de creadores...", 0.95)
                processed_rows, stats = analyze_duplicates(all_sw_results)
                self.analysis_rows = processed_rows
                self.stats_data = stats

                # 4. Actualizar tabla e interfaz
                self.after(0, lambda: self._update_results_ui())

            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error durante el Análisis", f"Ocurrió un error:\n{e}"))
            finally:
                self.after(0, lambda: self._reset_analysis_ui())

        threading.Thread(target=run_pipeline, daemon=True).start()

    def _reset_analysis_ui(self):
        self.btn_analyze.configure(state="normal")
        self.progress_bar.set(0)

    def _update_results_ui(self):
        # Actualizar tarjetas de estadísticas
        self.card_total_students.configure(text=str(self.stats_data.get("total_students", 0)))
        self.card_total_files.configure(text=str(self.stats_data.get("total_files", 0)))
        self.card_suspicious.configure(text=str(self.stats_data.get("suspicious_files", 0)))

        self._apply_table_filter()
        self.status_label.configure(text=f"Análisis completado. {self.stats_data.get('suspicious_files', 0)} pieza(s) con creador repetido entre distintos alumnos.")

    def _apply_table_filter(self):
        """Puebla la tabla filtrando según el checkbox 'Mostrar únicamente repetidos'."""
        for item in self.tree.get_children():
            self.tree.delete(item)

        only_duplicates = self.chk_only_duplicates.get()
        is_dark = ctk.get_appearance_mode() == "Dark"

        for row in self.analysis_rows:
            if only_duplicates and not row['is_duplicate']:
                continue

            tag = "ALERT_DARK" if is_dark and row['is_duplicate'] else ("ALERT" if row['is_duplicate'] else "NORMAL")
            
            self.tree.insert(
                "",
                "end",
                values=(
                    row['status'],
                    row['student_name'],
                    row['file_name'],
                    row['sw_author'],
                    row['sw_last_author'],
                    row['alert_details']
                ),
                tags=(tag,)
            )

    def _export_to_excel(self):
        """Exporta los resultados del análisis a un archivo Excel (.xlsx) o CSV."""
        if not self.analysis_rows:
            messagebox.showwarning("Sin Datos", "No hay datos de análisis para exportar.")
            return

        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx"), ("CSV Files", "*.csv")],
            title="Guardar Reporte de Autores de SolidWorks"
        )
        if not file_path:
            return

        try:
            if file_path.endswith('.xlsx'):
                import openpyxl
                from openpyxl.styles import PatternFill, Font

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Reporte Autores SW"

                headers = [
                    "Estado / Alerta", "Alumno (Classroom)", "Email Alumno",
                    "Pieza / Archivo", "Creador Registrado en SolidWorks",
                    "Último Cambio por (SW)", "Fecha Creación SW", "Detalles Coincidencias"
                ]
                ws.append(headers)

                # Estilo de encabezados
                header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                header_font = Font(color="FFFFFF", bold=True)
                for col_num in range(1, len(headers) + 1):
                    cell = ws.cell(row=1, column=col_num)
                    cell.fill = header_fill
                    cell.font = header_font

                red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

                for row_idx, r in enumerate(self.analysis_rows, start=2):
                    ws.append([
                        r['status'], r['student_name'], r['student_email'],
                        r['file_name'], r['sw_author'], r['sw_last_author'],
                        r['create_time'], r['alert_details']
                    ])

                    if r['is_duplicate']:
                        for c in range(1, len(headers) + 1):
                            ws.cell(row=row_idx, column=c).fill = red_fill

                wb.save(file_path)

            else:
                # Exportar CSV
                with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
                    writer = csv.writer(f)
                    writer.writerow([
                        "Estado", "Alumno", "Email", "Pieza", "Creador SolidWorks",
                        "Ultimo Cambio SW", "Fecha Creacion SW", "Detalles"
                    ])
                    for r in self.analysis_rows:
                        writer.writerow([
                            r['status'], r['student_name'], r['student_email'],
                            r['file_name'], r['sw_author'], r['sw_last_author'],
                            r['create_time'], r['alert_details']
                        ])

            messagebox.showinfo("Éxito", f"Reporte exportado correctamente a:\n{file_path}")
        except Exception as e:
            messagebox.showerror("Error al Exportar", f"No se pudo guardar el archivo:\n{e}")

if __name__ == "__main__":
    app = SolidWorksClassroomApp()
    app.mainloop()
